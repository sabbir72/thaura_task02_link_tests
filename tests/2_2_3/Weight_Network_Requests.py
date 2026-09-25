import argparse
from collections import Counter
from pathlib import Path
from urllib.parse import urlsplit

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.table import Table, TableStyleInfo

from playwright.sync_api import sync_playwright


# ============================================================
# TC-PERF-12
# Measure Total Page Weight and Network Requests
# ============================================================

PAGES = [
    "https://thaura.ai/home",
    "https://thaura.ai/story",
    "https://thaura.ai/constitution",
    "https://thaura.ai/pricing",
    "https://thaura.ai/community",
    "https://thaura.ai/contact",
    "https://thaura.ai/careers",
    "https://thaura.ai/faq",
    "https://thaura.ai/api-platform",
    "https://thaura.ai/donate",
    "https://thaura.ai/download",
    "https://thaura.ai/terms-of-service",
    "https://thaura.ai/privacy-policy",
    "https://thaura.ai/data-processing-addendum",
    "https://thaura.ai/imprint",
]

OUTPUT_DIR = Path("tc_output")
OUTPUT_DIR.mkdir(exist_ok=True)

OUTPUT_FILE = (
    OUTPUT_DIR
    / "Thaura_TC_Page_Weight_Network_Report.xlsx"
)


# ============================================================
# Helpers
# ============================================================

def get_kb(value):
    return round(value / 1024, 2)


def get_mb(value):
    return round(value / 1024 / 1024, 3)


def is_third_party(url):
    try:
        host = (
            urlsplit(url)
            .hostname
            or ""
        ).lower()

        return host not in {
            "thaura.ai",
            "www.thaura.ai",
        }

    except Exception:
        return False


def get_status_class(status):
    if 200 <= status < 300:
        return "2xx"

    if 300 <= status < 400:
        return "3xx"

    if 400 <= status < 500:
        return "4xx"

    if 500 <= status < 600:
        return "5xx"

    return "Other"


# ============================================================
# Collect one page
# ============================================================

def collect_page(page, page_url):

    requests_data = []
    failed_requests = []

    page_status = "N/A"
    final_url = page_url

    def on_request_finished(request):

        try:

            response = request.response()

            if response is None:
                return

            sizes = request.sizes()

            response_body_size = (
                sizes.get(
                    "responseBodySize",
                    0
                )
                or 0
            )

            request_headers_size = (
                sizes.get(
                    "requestHeadersSize",
                    0
                )
                or 0
            )

            response_headers_size = (
                sizes.get(
                    "responseHeadersSize",
                    0
                )
                or 0
            )

            request_url = request.url

            domain = (
                urlsplit(
                    request_url
                ).hostname
                or ""
            )

            status = response.status

            requests_data.append(
                {
                    "Page URL": page_url,

                    "Final Page URL": final_url,

                    "Request URL": request_url,

                    "Method": request.method,

                    "Resource Type":
                        request.resource_type,

                    "Status": status,

                    "Status Class":
                        get_status_class(
                            status
                        ),

                    "Content-Type":
                        response.header_value(
                            "content-type"
                        )
                        or "",

                    "Response Body Size (bytes)":
                        response_body_size,

                    "Response Body Size (KB)":
                        get_kb(
                            response_body_size
                        ),

                    "Response Body Size (MB)":
                        get_mb(
                            response_body_size
                        ),

                    "Request Headers Size (bytes)":
                        request_headers_size,

                    "Response Headers Size (bytes)":
                        response_headers_size,

                    "Domain": domain,

                    "Third Party":
                        "Yes"
                        if is_third_party(
                            request_url
                        )
                        else "No",
                }
            )

        except Exception as exc:

            print(
                "[WARNING] "
                f"requestfinished error: {exc}"
            )

    def on_request_failed(request):

        try:

            failed_requests.append(
                {
                    "Page URL": page_url,

                    "Request URL":
                        request.url,

                    "Method":
                        request.method,

                    "Resource Type":
                        request.resource_type,

                    "Failure":
                        request.failure
                        or "Unknown",

                    "Domain":
                        (
                            urlsplit(
                                request.url
                            ).hostname
                            or ""
                        ),

                    "Third Party":
                        "Yes"
                        if is_third_party(
                            request.url
                        )
                        else "No",
                }
            )

        except Exception as exc:

            print(
                "[WARNING] "
                f"requestfailed error: {exc}"
            )

    page.on(
        "requestfinished",
        on_request_finished
    )

    page.on(
        "requestfailed",
        on_request_failed
    )

    print("\n" + "=" * 110)

    print(
        f"TESTING: {page_url}"
    )

    print("=" * 110)

    try:

        response = page.goto(
            page_url,
            wait_until="load",
            timeout=60000
        )

        if response:

            page_status = response.status

        final_url = page.url

        # Wait for late resources.
        page.wait_for_timeout(3000)

        # Scroll to trigger lazy-loaded resources.
        try:

            page.evaluate(
                """
                async () => {

                    const wait = ms =>
                        new Promise(
                            resolve =>
                                setTimeout(
                                    resolve,
                                    ms
                                )
                        );

                    const maxHeight =
                        Math.max(
                            document.body.scrollHeight,
                            document.documentElement.scrollHeight
                        );

                    const step = 500;

                    for (
                        let y = 0;
                        y <= maxHeight;
                        y += step
                    ) {

                        window.scrollTo(
                            0,
                            y
                        );

                        await wait(200);
                    }

                    window.scrollTo(
                        0,
                        0
                    );

                    await wait(1500);
                }
                """
            )

        except Exception as exc:

            print(
                "[WARNING] "
                f"Scroll failed: {exc}"
            )

        # Give lazy-loaded requests time to finish.
        page.wait_for_timeout(2500)

    except Exception as exc:

        print(
            "[ERROR] "
            f"Page load failed: {exc}"
        )

    return (
        requests_data,
        failed_requests,
        page_status,
        final_url,
    )


# ============================================================
# Build Excel Workbook
# ============================================================

def build_excel(
    page_summaries,
    request_details,
    failed_requests,
    output_file,
):

    workbook = Workbook()

    # --------------------------------------------------------
    # Remove default sheet
    # --------------------------------------------------------

    default_sheet = workbook.active

    workbook.remove(
        default_sheet
    )

    # --------------------------------------------------------
    # Sheets
    # --------------------------------------------------------

    summary_ws = workbook.create_sheet(
        "Page Summary"
    )

    request_ws = workbook.create_sheet(
        "Request Details"
    )

    failed_ws = workbook.create_sheet(
        "Failed Requests"
    )

    # --------------------------------------------------------
    # Styles
    # --------------------------------------------------------

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="0F766E"
    )

    header_font = Font(
        bold=True,
        color="FFFFFF"
    )

    thin_border = Border(
        bottom=Side(
            style="thin",
            color="D1D5DB"
        )
    )

    header_alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

    body_alignment = Alignment(
        vertical="top",
        wrap_text=True
    )

    # --------------------------------------------------------
    # PAGE SUMMARY
    # --------------------------------------------------------

    summary_headers = [
        "Page URL",
        "Final Page URL",
        "Page HTTP Status",
        "Total Requests",
        "Total Transfer (bytes)",
        "Total Transfer (KB)",
        "Total Transfer (MB)",
        "Documents",
        "Scripts",
        "Stylesheets",
        "Images",
        "Fonts",
        "Media",
        "XHR/Fetch",
        "Other Requests",
        "Third Party Requests",
        "Third Party Scripts",
        "2xx",
        "3xx",
        "4xx",
        "5xx",
        "Failed Network Requests",
        "Largest Resource URL",
        "Largest Resource Type",
        "Largest Resource Size (KB)",
    ]

    summary_ws.append(
        summary_headers
    )

    for item in page_summaries:

        summary_ws.append(
            [
                item.get(
                    header,
                    ""
                )
                for header in summary_headers
            ]
        )

    # --------------------------------------------------------
    # REQUEST DETAILS
    # --------------------------------------------------------

    request_headers = [
        "Page URL",
        "Final Page URL",
        "Request URL",
        "Method",
        "Resource Type",
        "Status",
        "Status Class",
        "Content-Type",
        "Response Body Size (bytes)",
        "Response Body Size (KB)",
        "Response Body Size (MB)",
        "Request Headers Size (bytes)",
        "Response Headers Size (bytes)",
        "Domain",
        "Third Party",
    ]

    request_ws.append(
        request_headers
    )

    for item in request_details:

        request_ws.append(
            [
                item.get(
                    header,
                    ""
                )
                for header in request_headers
            ]
        )

    # --------------------------------------------------------
    # FAILED REQUESTS
    # --------------------------------------------------------

    failed_headers = [
        "Page URL",
        "Request URL",
        "Method",
        "Resource Type",
        "Failure",
        "Domain",
        "Third Party",
    ]

    failed_ws.append(
        failed_headers
    )

    for item in failed_requests:

        failed_ws.append(
            [
                item.get(
                    header,
                    ""
                )
                for header in failed_headers
            ]
        )

    # --------------------------------------------------------
    # Header formatting
    # --------------------------------------------------------

    for ws in [
        summary_ws,
        request_ws,
        failed_ws,
    ]:

        for cell in ws[1]:

            cell.fill = header_fill

            cell.font = header_font

            cell.alignment = (
                header_alignment
            )

            cell.border = thin_border

        ws.freeze_panes = "A2"

        ws.auto_filter.ref = (
            ws.dimensions
        )

        # Body alignment
        for row in ws.iter_rows(
            min_row=2
        ):

            for cell in row:

                cell.alignment = (
                    body_alignment
                )

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    # Summary
    summary_ws.column_dimensions[
        "A"
    ].width = 42

    summary_ws.column_dimensions[
        "B"
    ].width = 42

    summary_ws.column_dimensions[
        "W"
    ].width = 52

    for col in [
        "C",
        "D",
        "E",
        "F",
        "G",
        "H",
        "I",
        "J",
        "K",
        "L",
        "M",
        "N",
        "O",
        "P",
        "Q",
        "R",
        "S",
        "T",
        "U",
        "V",
        "X",
        "Y",
    ]:

        summary_ws.column_dimensions[
            col
        ].width = 16

    # Request details
    for col in [
        "A",
        "B",
        "C",
    ]:

        request_ws.column_dimensions[
            col
        ].width = 44

    for col in [
        "D",
        "E",
        "F",
        "G",
        "O",
    ]:

        request_ws.column_dimensions[
            col
        ].width = 18

    request_ws.column_dimensions[
        "H"
    ].width = 30

    request_ws.column_dimensions[
        "N"
    ].width = 28

    # Failed requests
    failed_ws.column_dimensions[
        "A"
    ].width = 42

    failed_ws.column_dimensions[
        "B"
    ].width = 55

    failed_ws.column_dimensions[
        "C"
    ].width = 15

    failed_ws.column_dimensions[
        "D"
    ].width = 18

    failed_ws.column_dimensions[
        "E"
    ].width = 30

    failed_ws.column_dimensions[
        "F"
    ].width = 25

    failed_ws.column_dimensions[
        "G"
    ].width = 15

    # --------------------------------------------------------
    # Number formats
    # --------------------------------------------------------

    for row in summary_ws.iter_rows(
        min_row=2,
        min_col=5,
        max_col=7
    ):

        for cell in row:
            cell.number_format = "0.00"

    for row in summary_ws.iter_rows(
        min_row=2,
        min_col=25,
        max_col=25
    ):

        for cell in row:
            cell.number_format = "0.00"

    for row in request_ws.iter_rows(
        min_row=2,
        min_col=9,
        max_col=11
    ):

        for cell in row:
            cell.number_format = "0.00"

    # --------------------------------------------------------
    # Error highlighting
    # --------------------------------------------------------

    if summary_ws.max_row >= 2:

        summary_ws.conditional_formatting.add(
            f"T2:U{summary_ws.max_row}",
            CellIsRule(
                operator="greaterThan",
                formula=["0"],
                fill=PatternFill(
                    fill_type="solid",
                    fgColor="FECACA"
                ),
                font=Font(
                    color="991B1B"
                ),
            )
        )

    if request_ws.max_row >= 2:

        request_ws.conditional_formatting.add(
            f"F2:F{request_ws.max_row}",
            CellIsRule(
                operator="greaterThanOrEqual",
                formula=["400"],
                fill=PatternFill(
                    fill_type="solid",
                    fgColor="FECACA"
                ),
                font=Font(
                    color="991B1B"
                ),
            )
        )

    # --------------------------------------------------------
    # Excel tables
    # --------------------------------------------------------

    if summary_ws.max_row >= 2:

        summary_table = Table(
            displayName="PageSummaryTable",
            ref=summary_ws.dimensions
        )

        summary_table.tableStyleInfo = (
            TableStyleInfo(
                name="TableStyleMedium4",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
        )

        summary_ws.add_table(
            summary_table
        )

    if request_ws.max_row >= 2:

        request_table = Table(
            displayName="RequestDetailsTable",
            ref=request_ws.dimensions
        )

        request_table.tableStyleInfo = (
            TableStyleInfo(
                name="TableStyleMedium4",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
        )

        request_ws.add_table(
            request_table
        )

    if failed_ws.max_row >= 2:

        failed_table = Table(
            displayName="FailedRequestsTable",
            ref=failed_ws.dimensions
        )

        failed_table.tableStyleInfo = (
            TableStyleInfo(
                name="TableStyleMedium4",
                showFirstColumn=False,
                showLastColumn=False,
                showRowStripes=True,
                showColumnStripes=False,
            )
        )

        failed_ws.add_table(
            failed_table
        )

    # --------------------------------------------------------
    # Workbook save
    # --------------------------------------------------------

    workbook.save(
        output_file
    )


# ============================================================
# Main
# ============================================================

def main():

    parser = argparse.ArgumentParser(
        description=(
            "TC-PERF-12 "
            "Page Weight and Network Requests"
        )
    )

    parser.add_argument(
        "-sv",
        "--show-browser",
        action="store_true",
        help=(
            "Run browser in headed mode"
        ),
    )

    args = parser.parse_args()

    page_summaries = []

    all_request_details = []

    all_failed_requests = []

    with sync_playwright() as p:

        browser = p.chromium.launch(
            headless=not args.show_browser
        )

        context = browser.new_context(
            viewport={
                "width": 1440,
                "height": 900,
            }
        )

        for page_url in PAGES:

            page = context.new_page()

            try:

                summary_requests = (
                    collect_page(
                        page,
                        page_url
                    )
                )

                (
                    requests_data,
                    failed_requests,
                    page_status,
                    final_url,
                ) = summary_requests

                # Resource counts
                resource_counter = Counter(
                    item[
                        "Resource Type"
                    ]
                    for item in requests_data
                )

                # HTTP status counts
                status_2xx = sum(
                    1
                    for item in requests_data
                    if 200
                    <= item["Status"]
                    < 300
                )

                status_3xx = sum(
                    1
                    for item in requests_data
                    if 300
                    <= item["Status"]
                    < 400
                )

                status_4xx = sum(
                    1
                    for item in requests_data
                    if 400
                    <= item["Status"]
                    < 500
                )

                status_5xx = sum(
                    1
                    for item in requests_data
                    if 500
                    <= item["Status"]
                    < 600
                )

                total_transfer = sum(
                    item[
                        "Response Body Size (bytes)"
                    ]
                    for item in requests_data
                )

                third_party_requests = [
                    item
                    for item in requests_data
                    if item[
                        "Third Party"
                    ] == "Yes"
                ]

                third_party_scripts = [
                    item
                    for item in third_party_requests
                    if item[
                        "Resource Type"
                    ] == "script"
                ]

                # Largest resources
                largest_resources = sorted(
                    requests_data,
                    key=lambda item:
                        item[
                            "Response Body Size (bytes)"
                        ],
                    reverse=True,
                )

                largest = (
                    largest_resources[0]
                    if largest_resources
                    else {}
                )

                summary = {

                    "Page URL":
                        page_url,

                    "Final Page URL":
                        final_url,

                    "Page HTTP Status":
                        page_status,

                    "Total Requests":
                        len(requests_data),

                    "Total Transfer (bytes)":
                        total_transfer,

                    "Total Transfer (KB)":
                        get_kb(
                            total_transfer
                        ),

                    "Total Transfer (MB)":
                        get_mb(
                            total_transfer
                        ),

                    "Documents":
                        resource_counter.get(
                            "document",
                            0
                        ),

                    "Scripts":
                        resource_counter.get(
                            "script",
                            0
                        ),

                    "Stylesheets":
                        resource_counter.get(
                            "stylesheet",
                            0
                        ),

                    "Images":
                        resource_counter.get(
                            "image",
                            0
                        ),

                    "Fonts":
                        resource_counter.get(
                            "font",
                            0
                        ),

                    "Media":
                        resource_counter.get(
                            "media",
                            0
                        ),

                    "XHR/Fetch":
                        (
                            resource_counter.get(
                                "xhr",
                                0
                            )
                            +
                            resource_counter.get(
                                "fetch",
                                0
                            )
                        ),

                    "Other Requests":
                        (
                            len(requests_data)
                            -
                            sum(
                                resource_counter.get(
                                    key,
                                    0
                                )
                                for key in [
                                    "document",
                                    "script",
                                    "stylesheet",
                                    "image",
                                    "font",
                                    "media",
                                    "xhr",
                                    "fetch",
                                ]
                            )
                        ),

                    "Third Party Requests":
                        len(
                            third_party_requests
                        ),

                    "Third Party Scripts":
                        len(
                            third_party_scripts
                        ),

                    "2xx":
                        status_2xx,

                    "3xx":
                        status_3xx,

                    "4xx":
                        status_4xx,

                    "5xx":
                        status_5xx,

                    "Failed Network Requests":
                        len(
                            failed_requests
                        ),

                    "Largest Resource URL":
                        largest.get(
                            "Request URL",
                            ""
                        ),

                    "Largest Resource Type":
                        largest.get(
                            "Resource Type",
                            ""
                        ),

                    "Largest Resource Size (KB)":
                        largest.get(
                            "Response Body Size (KB)",
                            0
                        ),
                }

                page_summaries.append(
                    summary
                )

                all_request_details.extend(
                    requests_data
                )

                all_failed_requests.extend(
                    failed_requests
                )

                print(
                    f"[RESULT] "
                    f"requests="
                    f"{len(requests_data)} | "
                    f"transfer="
                    f"{get_mb(total_transfer)} MB | "
                    f"3xx={status_3xx} | "
                    f"4xx={status_4xx} | "
                    f"5xx={status_5xx} | "
                    f"failed="
                    f"{len(failed_requests)}"
                )

            finally:

                page.close()

        context.close()

        browser.close()

    # --------------------------------------------------------
    # Create Excel
    # --------------------------------------------------------

    build_excel(
        page_summaries,
        all_request_details,
        all_failed_requests,
        OUTPUT_FILE,
    )

    # --------------------------------------------------------
    # Final output
    # --------------------------------------------------------

    print("\n" + "=" * 110)

    print(
        "TC- TEST COMPLETED"
    )

    print("=" * 110)

    print(
        f"Pages tested: "
        f"{len(PAGES)}"
    )

    print(
        f"Total requests collected: "
        f"{len(all_request_details)}"
    )

    print(
        f"Failed network requests: "
        f"{len(all_failed_requests)}"
    )

    print(
        f"Excel report: "
        f"{OUTPUT_FILE.resolve()}"
    )


if __name__ == "__main__":
    main()