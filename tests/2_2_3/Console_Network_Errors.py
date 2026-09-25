import argparse
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.worksheet.table import (
    Table,
    TableStyleInfo,
)

from playwright.sync_api import sync_playwright


# ============================================================
# TC-PERF-13
# Console & Network Errors
# ============================================================

PAGES = [
    "https://thaura.ai/home",
    "https://thaura.ai/story",
    "https://thaura.ai/constitution",
    "https://thaura.ai/pricing",
    "https://thaura.ai/community",
    "https://thaura.ai/contact",
    "https://thaura.ai/careers",
    "https://thaura.ai/faq",
    "https://thaura.ai/api-platform",
    "https://thaura.ai/donate",
    "https://thaura.ai/download",
    "https://thaura.ai/terms-of-service",
    "https://thaura.ai/privacy-policy",
    "https://thaura.ai/data-processing-addendum",
    "https://thaura.ai/imprint",
]


OUTPUT_DIR = Path(
    "tc_perf_13_output"
)

OUTPUT_DIR.mkdir(
    exist_ok=True
)


OUTPUT_XLSX = (
    OUTPUT_DIR
    / "Thaura_TC_Console_Network_Errors.xlsx"
)


SCREENSHOT_DIR = (
    OUTPUT_DIR
    / "screenshots"
)


SCREENSHOT_DIR.mkdir(
    exist_ok=True
)


# ============================================================
# Helpers
# ============================================================

def safe_text(value):
    if value is None:
        return ""

    return str(value)


def get_status_class(status):

    if 200 <= status < 300:
        return "2xx"

    if 300 <= status < 400:
        return "3xx"

    if 400 <= status < 500:
        return "4xx"

    if 500 <= status < 600:
        return "5xx"

    return "Other"


def safe_console_location(message):

    try:

        location = message.location

        if callable(location):
            location = location()

        if not isinstance(
            location,
            dict
        ):
            return "", "", ""

        return (
            safe_text(
                location.get("url")
            ),
            location.get(
                "line",
                ""
            ),
            location.get(
                "column",
                ""
            ),
        )

    except Exception:

        return "", "", ""


def slug_from_url(url):

    path = re.sub(
        r"[^a-zA-Z0-9]+",
        "_",
        url.rstrip("/")
        .split("//", 1)[-1]
    )

    return (
        path.strip("_")[:100]
        or "page"
    )


# ============================================================
# Test One Page
# ============================================================

def collect_page(
    context,
    page_url
):

    page = context.new_page()

    console_errors = []

    console_warnings = []

    page_errors = []

    http_errors = []

    network_failures = []

    page_status = "N/A"

    final_url = page_url

    navigation_error = ""


    # ========================================================
    # Console Messages
    # ========================================================

    def on_console(message):

        try:

            message_type = (
                safe_text(
                    message.type
                )
                .lower()
            )

            if message_type not in {
                "error",
                "warning",
            }:
                return


            source_url, line, column = (
                safe_console_location(
                    message
                )
            )


            record = {

                "Page URL":
                    page_url,

                "Console Type":
                    message_type,

                "Message":
                    safe_text(
                        message.text
                    ),

                "Source URL":
                    source_url,

                "Line":
                    line,

                "Column":
                    column,
            }


            if message_type == "error":

                console_errors.append(
                    record
                )

            elif message_type == "warning":

                console_warnings.append(
                    record
                )


        except Exception as exc:

            print(
                "[WARNING] "
                f"Console handler error: "
                f"{exc}"
            )


    page.on(
        "console",
        on_console
    )


    # ========================================================
    # Uncaught Page Errors
    # ========================================================

    def on_page_error(error):

        try:

            message = getattr(
                error,
                "message",
                None
            )

            if callable(message):
                message = message()


            page_errors.append(
                {

                    "Page URL":
                        page_url,

                    "Error Message":
                        safe_text(
                            message
                            or error
                        ),
                }
            )


        except Exception as exc:

            print(
                "[WARNING] "
                f"Page error handler: "
                f"{exc}"
            )


    page.on(
        "pageerror",
        on_page_error
    )


    # ========================================================
    # HTTP Response Monitoring
    # ========================================================

    def on_response(response):

        try:

            status = response.status

            # Only record HTTP errors.
            if status < 400:
                return


            request = response.request


            http_errors.append(
                {

                    "Page URL":
                        page_url,

                    "Request URL":
                        response.url,

                    "Method":
                        request.method,

                    "Resource Type":
                        request.resource_type,

                    "Status":
                        status,

                    "Status Class":
                        get_status_class(
                            status
                        ),

                    "Status Text":
                        safe_text(
                            response.status_text
                        ),

                    "Content-Type":
                        (
                            response.header_value(
                                "content-type"
                            )
                            or ""
                        ),
                }
            )


        except Exception as exc:

            print(
                "[WARNING] "
                f"Response handler error: "
                f"{exc}"
            )


    page.on(
        "response",
        on_response
    )


    # ========================================================
    # Network Failures
    # ========================================================

    def on_request_failed(request):

        try:

            failure = request.failure

            if callable(failure):
                failure = failure()


            network_failures.append(
                {

                    "Page URL":
                        page_url,

                    "Request URL":
                        request.url,

                    "Method":
                        request.method,

                    "Resource Type":
                        request.resource_type,

                    "Failure":
                        safe_text(
                            failure
                            or "Unknown"
                        ),
                }
            )


        except Exception as exc:

            print(
                "[WARNING] "
                f"Network failure handler: "
                f"{exc}"
            )


    page.on(
        "requestfailed",
        on_request_failed
    )


    # ========================================================
    # Execute Page
    # ========================================================

    print("\n" + "=" * 110)

    print(
        f"TESTING: {page_url}"
    )

    print("=" * 110)


    try:

        response = page.goto(
            page_url,
            wait_until="load",
            timeout=60000
        )


        if response is not None:

            page_status = (
                response.status
            )


        final_url = page.url


        # Wait for JavaScript and
        # asynchronous resources.
        page.wait_for_timeout(
            3000
        )


        # ====================================================
        # Scroll page
        # ====================================================

        try:

            page.evaluate(
                """
                async () => {

                    const wait = ms =>
                        new Promise(
                            resolve =>
                                setTimeout(
                                    resolve,
                                    ms
                                )
                        );

                    const maxHeight =
                        Math.max(
                            document.body.scrollHeight,
                            document.documentElement.scrollHeight
                        );

                    const step = 500;

                    for (
                        let y = 0;
                        y <= maxHeight;
                        y += step
                    ) {

                        window.scrollTo(
                            0,
                            y
                        );

                        await wait(200);
                    }

                    window.scrollTo(
                        0,
                        0
                    );

                    await wait(1500);
                }
                """
            )


        except Exception as exc:

            print(
                "[WARNING] "
                f"Scroll failed: "
                f"{exc}"
            )


        # Wait for additional
        # requests/errors.
        page.wait_for_timeout(
            2500
        )


    except Exception as exc:

        navigation_error = (
            safe_text(exc)
        )

        print(
            "[ERROR] "
            f"Page execution: "
            f"{navigation_error}"
        )


    # ========================================================
    # Final Page Status
    # ========================================================

    if (
        console_errors
        or page_errors
        or http_errors
        or network_failures
        or navigation_error
    ):

        page_final_status = "FAIL"

    elif console_warnings:

        page_final_status = (
            "PASS WITH WARNING"
        )

    else:

        page_final_status = "PASS"


    # ========================================================
    # Screenshot
    # ========================================================

    screenshot_path = ""


    if (
        console_errors
        or console_warnings
        or page_errors
        or http_errors
        or network_failures
        or navigation_error
    ):

        screenshot_path = str(

            SCREENSHOT_DIR
            / (
                f"{slug_from_url(page_url)}"
                ".png"
            )
        )


        try:

            page.screenshot(
                path=screenshot_path,
                full_page=True,
            )


        except Exception as exc:

            print(
                "[WARNING] "
                f"Screenshot failed: "
                f"{exc}"
            )

            screenshot_path = ""


    # ========================================================
    # Summary
    # ========================================================

    summary = {

        "Page URL":
            page_url,

        "Final Page URL":
            final_url,

        "Page HTTP Status":
            page_status,

        "Console Errors":
            len(
                console_errors
            ),

        "Console Warnings":
            len(
                console_warnings
            ),

        "Page Errors":
            len(
                page_errors
            ),

        "HTTP 4xx":
            sum(
                1
                for item in http_errors
                if item["Status Class"]
                == "4xx"
            ),

        "HTTP 5xx":
            sum(
                1
                for item in http_errors
                if item["Status Class"]
                == "5xx"
            ),

        "Network Failures":
            len(
                network_failures
            ),

        "Navigation Error":
            navigation_error,

        "Final Status":
            page_final_status,

        "Evidence Screenshot":
            screenshot_path,
    }


    print(
        f"[RESULT] "
        f"console_error="
        f"{len(console_errors)} | "
        f"console_warning="
        f"{len(console_warnings)} | "
        f"page_error="
        f"{len(page_errors)} | "
        f"4xx="
        f"{summary['HTTP 4xx']} | "
        f"5xx="
        f"{summary['HTTP 5xx']} | "
        f"network_failure="
        f"{len(network_failures)} | "
        f"status="
        f"{page_final_status}"
    )


    page.close()


    return (

        summary,

        console_errors,

        console_warnings,

        page_errors,

        http_errors,

        network_failures,
    )


# ============================================================
# Excel Styling
# ============================================================

def style_sheet(
    ws,
    widths
):

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="0F766E",
    )


    header_font = Font(
        bold=True,
        color="FFFFFF",
    )


    header_alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True,
    )


    body_alignment = Alignment(
        vertical="top",
        wrap_text=True,
    )


    border = Border(
        bottom=Side(
            style="thin",
            color="D1D5DB",
        )
    )


    # Header
    for cell in ws[1]:

        cell.fill = header_fill

        cell.font = header_font

        cell.alignment = (
            header_alignment
        )

        cell.border = border


    # Body
    for row in ws.iter_rows(
        min_row=2
    ):

        for cell in row:

            cell.alignment = (
                body_alignment
            )


    ws.freeze_panes = "A2"

    ws.auto_filter.ref = (
        ws.dimensions
    )

    ws.row_dimensions[1].height = 32


    # Widths
    for column, width in (
        widths.items()
    ):

        ws.column_dimensions[
            column
        ].width = width


def add_table(
    ws,
    table_name
):

    if ws.max_row < 2:
        return


    table = Table(
        displayName=table_name,
        ref=ws.dimensions,
    )


    table.tableStyleInfo = (
        TableStyleInfo(

            name="TableStyleMedium4",

            showFirstColumn=False,

            showLastColumn=False,

            showRowStripes=True,

            showColumnStripes=False,
        )
    )


    ws.add_table(
        table
    )


# ============================================================
# Excel Report
# ============================================================

def build_excel(
    summaries,
    console_errors,
    console_warnings,
    page_errors,
    http_errors,
    network_failures,
):

    workbook = Workbook()


    # ========================================================
    # Sheets
    # ========================================================

    summary_ws = workbook.active

    summary_ws.title = (
        "Summary"
    )


    console_error_ws = (
        workbook.create_sheet(
            "Console Errors"
        )
    )


    console_warning_ws = (
        workbook.create_sheet(
            "Console Warnings"
        )
    )


    page_error_ws = (
        workbook.create_sheet(
            "Page Errors"
        )
    )


    http_error_ws = (
        workbook.create_sheet(
            "HTTP Errors"
        )
    )


    network_failure_ws = (
        workbook.create_sheet(
            "Network Failures"
        )
    )


    # ========================================================
    # Summary Sheet
    # ========================================================

    summary_headers = [

        "Page URL",

        "Final Page URL",

        "Page HTTP Status",

        "Console Errors",

        "Console Warnings",

        "Page Errors",

        "HTTP 4xx",

        "HTTP 5xx",

        "Network Failures",

        "Navigation Error",

        "Final Status",

        "Evidence Screenshot",
    ]


    summary_ws.append(
        summary_headers
    )


    for item in summaries:

        summary_ws.append(
            [
                item.get(
                    header,
                    ""
                )
                for header
                in summary_headers
            ]
        )


    # ========================================================
    # Console Errors
    # ========================================================

    console_error_headers = [

        "Page URL",

        "Console Type",

        "Message",

        "Source URL",

        "Line",

        "Column",
    ]


    console_error_ws.append(
        console_error_headers
    )


    for item in console_errors:

        console_error_ws.append(
            [
                item.get(
                    header,
                    ""
                )
                for header
                in console_error_headers
            ]
        )


    # ========================================================
    # Console Warnings
    # ========================================================

    console_warning_headers = [

        "Page URL",

        "Console Type",

        "Message",

        "Source URL",

        "Line",

        "Column",
    ]


    console_warning_ws.append(
        console_warning_headers
    )


    for item in console_warnings:

        console_warning_ws.append(
            [
                item.get(
                    header,
                    ""
                )
                for header
                in console_warning_headers
            ]
        )


    # ========================================================
    # Page Errors
    # ========================================================

    page_error_headers = [

        "Page URL",

        "Error Message",
    ]


    page_error_ws.append(
        page_error_headers
    )


    for item in page_errors:

        page_error_ws.append(
            [
                item.get(
                    header,
                    ""
                )
                for header
                in page_error_headers
            ]
        )


    # ========================================================
    # HTTP Errors
    # ========================================================

    http_error_headers = [

        "Page URL",

        "Request URL",

        "Method",

        "Resource Type",

        "Status",

        "Status Class",

        "Status Text",

        "Content-Type",
    ]


    http_error_ws.append(
        http_error_headers
    )


    for item in http_errors:

        http_error_ws.append(
            [
                item.get(
                    header,
                    ""
                )
                for header
                in http_error_headers
            ]
        )


    # ========================================================
    # Network Failures
    # ========================================================

    network_failure_headers = [

        "Page URL",

        "Request URL",

        "Method",

        "Resource Type",

        "Failure",
    ]


    network_failure_ws.append(
        network_failure_headers
    )


    for item in network_failures:

        network_failure_ws.append(
            [
                item.get(
                    header,
                    ""
                )
                for header
                in network_failure_headers
            ]
        )


    # ========================================================
    # Styling
    # ========================================================

    style_sheet(
        summary_ws,
        {
            "A": 42,
            "B": 42,
            "C": 16,
            "D": 16,
            "E": 17,
            "F": 14,
            "G": 12,
            "H": 12,
            "I": 18,
            "J": 45,
            "K": 22,
            "L": 55,
        },
    )


    style_sheet(
        console_error_ws,
        {
            "A": 42,
            "B": 18,
            "C": 80,
            "D": 55,
            "E": 12,
            "F": 12,
        },
    )


    style_sheet(
        console_warning_ws,
        {
            "A": 42,
            "B": 18,
            "C": 80,
            "D": 55,
            "E": 12,
            "F": 12,
        },
    )


    style_sheet(
        page_error_ws,
        {
            "A": 42,
            "B": 100,
        },
    )


    style_sheet(
        http_error_ws,
        {
            "A": 42,
            "B": 70,
            "C": 12,
            "D": 18,
            "E": 12,
            "F": 14,
            "G": 20,
            "H": 30,
        },
    )


    style_sheet(
        network_failure_ws,
        {
            "A": 42,
            "B": 70,
            "C": 12,
            "D": 18,
            "E": 45,
        },
    )


    # ========================================================
    # Conditional Formatting
    # ========================================================

    if summary_ws.max_row >= 2:

        summary_ws.conditional_formatting.add(

            f"D2:I{summary_ws.max_row}",

            CellIsRule(

                operator="greaterThan",

                formula=["0"],

                fill=PatternFill(
                    fill_type="solid",
                    fgColor="FECACA",
                ),

                font=Font(
                    color="991B1B"
                ),
            )
        )


    # ========================================================
    # Tables
    # ========================================================

    add_table(
        summary_ws,
        "SummaryTable"
    )


    add_table(
        console_error_ws,
        "ConsoleErrorsTable"
    )


    add_table(
        console_warning_ws,
        "ConsoleWarningsTable"
    )


    add_table(
        page_error_ws,
        "PageErrorsTable"
    )


    add_table(
        http_error_ws,
        "HTTPErrorsTable"
    )


    add_table(
        network_failure_ws,
        "NetworkFailuresTable"
    )


    # ========================================================
    # Save
    # ========================================================

    workbook.save(
        OUTPUT_XLSX
    )


# ============================================================
# Main
# ============================================================

def main():

    parser = argparse.ArgumentParser(

        description=(
            "TC-PERF-13 - "
            "Console & Network Error Monitor"
        )
    )


    parser.add_argument(

        "-sv",

        "--show-browser",

        action="store_true",

        help=(
            "Run Chromium in headed mode."
        ),
    )


    args = parser.parse_args()


    summaries = []

    all_console_errors = []

    all_console_warnings = []

    all_page_errors = []

    all_http_errors = []

    all_network_failures = []


    # ========================================================
    # Playwright
    # ========================================================

    with sync_playwright() as p:

        browser = p.chromium.launch(

            headless=not args.show_browser
        )


        context = browser.new_context(

            viewport={
                "width": 1440,
                "height": 900,
            }
        )


        # ====================================================
        # Test All 15 Pages
        # ====================================================

        for page_url in PAGES:

            (
                summary,

                console_errors,

                console_warnings,

                page_errors,

                http_errors,

                network_failures,

            ) = collect_page(

                context,

                page_url
            )


            summaries.append(
                summary
            )


            all_console_errors.extend(
                console_errors
            )


            all_console_warnings.extend(
                console_warnings
            )


            all_page_errors.extend(
                page_errors
            )


            all_http_errors.extend(
                http_errors
            )


            all_network_failures.extend(
                network_failures
            )


        context.close()

        browser.close()


    # ========================================================
    # Build Excel
    # ========================================================

    build_excel(

        summaries,

        all_console_errors,

        all_console_warnings,

        all_page_errors,

        all_http_errors,

        all_network_failures,
    )


    # ========================================================
    # Overall Final Status
    # ========================================================

    if (

        all_console_errors

        or all_page_errors

        or all_http_errors

        or all_network_failures

    ):

        overall_status = (
            "FAIL"
        )

    elif all_console_warnings:

        overall_status = (
            "PASS WITH WARNING"
        )

    else:

        overall_status = (
            "PASS"
        )


    # ========================================================
    # Totals
    # ========================================================

    total_4xx = sum(

        item["HTTP 4xx"]

        for item
        in summaries
    )


    total_5xx = sum(

        item["HTTP 5xx"]

        for item
        in summaries
    )


    total_network_failures = sum(

        item["Network Failures"]

        for item
        in summaries
    )


    # ========================================================
    # Final Console Output
    # ========================================================

    print("\n")

    print("=" * 110)

    print(
        "TC-PERF-13 TEST COMPLETED"
    )

    print("=" * 110)


    print(
        f"Pages tested       : "
        f"{len(PAGES)}"
    )


    print(
        f"Console errors     : "
        f"{len(all_console_errors)}"
    )


    print(
        f"Console warnings   : "
        f"{len(all_console_warnings)}"
    )


    print(
        f"Page errors        : "
        f"{len(all_page_errors)}"
    )


    print(
        f"HTTP 4xx           : "
        f"{total_4xx}"
    )


    print(
        f"HTTP 5xx           : "
        f"{total_5xx}"
    )


    print(
        f"Network failures   : "
        f"{total_network_failures}"
    )


    print(
        f"FINAL STATUS       : "
        f"{overall_status}"
    )


    print(
        f"Excel report       : "
        f"{OUTPUT_XLSX.resolve()}"
    )


if __name__ == "__main__":

    main()